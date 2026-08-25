#!/usr/bin/env python3
"""Export dissertation-appendix tables from data/sets.json.

Produces:
- CSV (set-level and superimposition-level)
- LaTeX (single-part and two-part)
- Word-friendly .doc (HTML content; single-part and two-part)
- Native .docx (single-part and two-part)
- Excel-friendly SpreadsheetML XML (single-part and two-part)
"""

import csv
import json
import re
from pathlib import Path

try:
    from docx import Document
    from docx.shared import Pt, RGBColor, Cm
    from docx.oxml.ns import qn
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False

ROOT = Path(__file__).resolve().parents[1]
DATA_PATH = ROOT / "data" / "sets.json"
EXPORT_DIR = ROOT / "exports"

HEXATONIC_PARENT_LABELS = [
    {"name": "HEX 0,1 Pentatonic Subset", "pcs": [0, 1, 4, 5, 8, 9]},
    {"name": "HEX 3,4 Pentatonic Subset", "pcs": [0, 3, 4, 7, 8, 11]},
]

DEGREE_TO_PITCH_CLASS = {
    "1": 0,
    "b2": 1,
    "b9": 1,
    "2": 2,
    "9": 2,
    "#2": 3,
    "#9": 3,
    "b3": 3,
    "3": 4,
    "b4": 4,
    "4": 5,
    "#4": 6,
    "#11": 6,
    "b5": 6,
    "5": 7,
    "#5": 8,
    "b6": 8,
    "b13": 8,
    "6": 9,
    "13": 9,
    "bb7": 9,
    "b7": 10,
    "7": 11,
    "maj7": 11,
    "addmaj7": 11,
}

COLLECTION_FAMILIES = [
    {
        "family": "Major Modes",
        "modes": [
            {"name": "Major", "degrees": ["1", "2", "3", "4", "5", "6", "7"]},
            {"name": "Dorian", "degrees": ["1", "2", "b3", "4", "5", "6", "b7"]},
            {"name": "Phrygian", "degrees": ["1", "b2", "b3", "4", "5", "b6", "b7"]},
            {"name": "Lydian", "degrees": ["1", "2", "3", "#4", "5", "6", "7"]},
            {"name": "Mixolydian", "degrees": ["1", "2", "3", "4", "5", "6", "b7"]},
            {"name": "Aeolian", "degrees": ["1", "2", "b3", "4", "5", "b6", "b7"]},
            {"name": "Locrian", "degrees": ["1", "b2", "b3", "4", "b5", "b6", "b7"]},
        ],
    },
    {
        "family": "Melodic Minor Modes",
        "modes": [
            {"name": "Melodic Minor", "degrees": ["1", "2", "b3", "4", "5", "6", "7"]},
            {"name": "Dorian b2", "degrees": ["1", "b2", "b3", "4", "5", "6", "b7"]},
            {"name": "Lydian Augmented", "degrees": ["1", "2", "3", "#4", "#5", "6", "7"]},
            {"name": "Lydian Dominant", "degrees": ["1", "2", "3", "#4", "5", "6", "b7"]},
            {"name": "Mixolydian b6", "degrees": ["1", "2", "3", "4", "5", "b6", "b7"]},
            {"name": "Aeolian b5", "degrees": ["1", "2", "b3", "4", "b5", "b6", "b7"]},
            {"name": "Altered Dominant", "degrees": ["1", "b2", "b3", "3", "b5", "b6", "b7"]},
        ],
    },
    {
        "family": "Harmonic Minor Modes",
        "modes": [
            {"name": "Harmonic Minor", "degrees": ["1", "2", "b3", "4", "5", "b6", "7"]},
            {"name": "Locrian Maj6", "degrees": ["1", "b2", "b3", "4", "b5", "6", "b7"]},
            {"name": "Ionian Augmented", "degrees": ["1", "2", "3", "4", "#5", "6", "7"]},
            {"name": "Dorian #11", "degrees": ["1", "2", "b3", "#4", "5", "6", "b7"]},
            {"name": "Phrygian Dominant", "degrees": ["1", "b2", "3", "4", "5", "b6", "b7"]},
            {"name": "Lydian #2", "degrees": ["1", "#2", "3", "#4", "5", "6", "7"]},
            {"name": "Super Locrian bb7", "degrees": ["1", "b2", "b3", "3", "b5", "b6", "bb7"]},
        ],
    },
    {
        "family": "Harmonic Major Modes",
        "modes": [
            {"name": "Harmonic Major", "degrees": ["1", "2", "3", "4", "5", "b6", "7"]},
            {"name": "Dorian b5", "degrees": ["1", "2", "b3", "4", "b5", "6", "b7"]},
            {"name": "Phrygian b4", "degrees": ["1", "b2", "b3", "b4", "5", "b6", "b7"]},
            {"name": "Lydian Minor", "degrees": ["1", "2", "b3", "#4", "5", "6", "7"]},
            {"name": "Mixolydian b2", "degrees": ["1", "b2", "3", "4", "5", "6", "b7"]},
            {"name": "Lydian Augmented #2", "degrees": ["1", "#2", "3", "#4", "#5", "6", "7"]},
            {"name": "Locrian bb7", "degrees": ["1", "b2", "b3", "4", "b5", "b6", "bb7"]},
        ],
    },
]

SPECIAL_COLLECTIONS = [
    {"name": "Augmented Scale", "degrees": ["1", "b3", "3", "5", "b6", "7"]},
    {"name": "Augmented Scale Half-Third", "degrees": ["1", "b2", "3", "4", "#5", "6"]},
]


def join(values):
    return ", ".join(str(v) for v in values)


def get_alteration_count(refs):
    for ref in refs or []:
        m = re.search(r"(\d+)\s+alteration", ref)
        if m:
            return int(m.group(1))
    return 0


def get_hexatonic_parent_labels(item):
    pcs = set(item.get("pcs_transposed_to_0") or item.get("pcs") or [])
    labels = []
    for collection in HEXATONIC_PARENT_LABELS:
        if pcs and pcs.issubset(set(collection["pcs"])):
            labels.append(collection["name"])
    return labels


def get_parent_pentatonic_labels(item):
    refs = [
        r
        for r in (item.get("pentatonic_reference") or [])
        if not ("Major Pentatonic" in r and "alteration" in r)
    ]
    merged = refs + get_hexatonic_parent_labels(item)
    deduped = []
    for label in merged:
        if label not in deduped:
            deduped.append(label)
    return deduped


def get_pitch_class_set(degrees):
    out = set()
    for degree in degrees or []:
        pc = DEGREE_TO_PITCH_CLASS.get(degree)
        if pc is not None:
            out.add(pc)
    return out


def is_subset_of_scale(chord_tones, scale_degrees):
    chord_set = get_pitch_class_set(chord_tones)
    scale_set = get_pitch_class_set(scale_degrees)
    return bool(chord_set) and chord_set.issubset(scale_set)


def get_modal_collection_labels(chord_tones):
    labels = []
    for family in COLLECTION_FAMILIES:
        for mode in family["modes"]:
            if is_subset_of_scale(chord_tones, mode["degrees"]):
                labels.append(mode["name"])
    for collection in SPECIAL_COLLECTIONS:
        if is_subset_of_scale(chord_tones, collection["degrees"]):
            labels.append(collection["name"])
    return labels


def get_voicing_collections(voicing):
    base = voicing.get("superset_collections") or []
    derived = get_modal_collection_labels(voicing.get("chord_tones") or [])
    deduped = []
    for c in base + derived:
        if c not in deduped:
            deduped.append(c)
    return deduped


def to_tex(text):
    return (
        str(text)
        .replace("\\", "\\textbackslash{}")
        .replace("&", "\\&")
        .replace("%", "\\%")
        .replace("$", "\\$")
        .replace("#", "\\#")
        .replace("_", "\\_")
        .replace("{", "\\{")
        .replace("}", "\\}")
        .replace("~", "\\textasciitilde{}")
        .replace("^", "\\textasciicircum{}")
    )


def to_xml(text):
    return (
        str(text)
        .replace("&", "&amp;")
        .replace("<", "&lt;")
        .replace(">", "&gt;")
        .replace('"', "&quot;")
        .replace("'", "&apos;")
    )


def write_word_doc(path, title, headers, rows):
    # Word opens HTML content with .doc extension reliably for appendix workflows.
    parts = [
        "<html><head><meta charset='utf-8'>",
        "<style>body{font-family:Times New Roman,serif;font-size:11pt;}"
        "h1{font-size:14pt;margin-bottom:8pt;}"
        "table{border-collapse:collapse;width:100%;}"
        "th,td{border:1px solid #222;padding:4px;vertical-align:top;}"
        "th{background:#efefef;}</style></head><body>",
        f"<h1>{to_xml(title)}</h1>",
        "<table><thead><tr>",
    ]
    for h in headers:
        parts.append(f"<th>{to_xml(h)}</th>")
    parts.append("</tr></thead><tbody>")
    for row in rows:
        parts.append("<tr>")
        for value in row:
            parts.append(f"<td>{to_xml(value)}</td>")
        parts.append("</tr>")
    parts.append("</tbody></table></body></html>")
    path.write_text("".join(parts), encoding="utf-8")


def write_word_two_part_doc(path, title, sets_headers, sets_rows, voicing_headers, voicing_rows):
    parts = [
        "<html><head><meta charset='utf-8'>",
        "<style>body{font-family:Times New Roman,serif;font-size:11pt;}"
        "h1{font-size:14pt;margin:0 0 10pt 0;}"
        "h2{font-size:12pt;margin:14pt 0 6pt 0;}"
        "table{border-collapse:collapse;width:100%;margin-bottom:12pt;}"
        "th,td{border:1px solid #222;padding:4px;vertical-align:top;}"
        "th{background:#efefef;}</style></head><body>",
        f"<h1>{to_xml(title)}</h1>",
        "<h2>Part A: Set-Level Table</h2>",
        "<table><thead><tr>",
    ]
    for h in sets_headers:
        parts.append(f"<th>{to_xml(h)}</th>")
    parts.append("</tr></thead><tbody>")
    for row in sets_rows:
        parts.append("<tr>")
        for value in row:
            parts.append(f"<td>{to_xml(value)}</td>")
        parts.append("</tr>")
    parts.append("</tbody></table>")

    parts.append("<h2>Part B: Superimposition Table</h2>")
    parts.append("<table><thead><tr>")
    for h in voicing_headers:
        parts.append(f"<th>{to_xml(h)}</th>")
    parts.append("</tr></thead><tbody>")
    for row in voicing_rows:
        parts.append("<tr>")
        for value in row:
            parts.append(f"<td>{to_xml(value)}</td>")
        parts.append("</tr>")
    parts.append("</tbody></table></body></html>")
    path.write_text("".join(parts), encoding="utf-8")


def _docx_header_row(table, headers):
    """Write a bold, shaded header row into a python-docx Table."""
    hdr_row = table.rows[0]
    for idx, text in enumerate(headers):
        cell = hdr_row.cells[idx]
        cell.text = text
        run = cell.paragraphs[0].runs[0]
        run.bold = True
        run.font.size = Pt(9)
        shading = cell._tc.get_or_add_tcPr()
        shading_el = shading.find(qn('w:shd'))
        if shading_el is None:
            from lxml import etree
            shading_el = etree.SubElement(shading, qn('w:shd'))
        shading_el.set(qn('w:val'), 'clear')
        shading_el.set(qn('w:color'), 'auto')
        shading_el.set(qn('w:fill'), 'EFEFEF')


def _docx_data_rows(table, rows, start_row=1):
    """Populate data rows in a python-docx Table starting at start_row."""
    for r_idx, row in enumerate(rows, start_row):
        if r_idx < len(table.rows):
            tr = table.rows[r_idx]
        else:
            tr = table.add_row()
        for c_idx, value in enumerate(row):
            cell = tr.cells[c_idx]
            cell.text = str(value)
            run = cell.paragraphs[0].runs[0] if cell.paragraphs[0].runs else cell.paragraphs[0].add_run(str(value))
            if not cell.paragraphs[0].runs:
                run.text = str(value)
            run.font.size = Pt(8)


def _set_table_style(table):
    table.style = 'Table Grid'


def write_docx_single(path, title, headers, rows):
    if not DOCX_AVAILABLE:
        return
    doc = Document()
    doc.core_properties.title = title
    heading = doc.add_heading(title, level=1)
    heading.runs[0].font.size = Pt(14)
    table = doc.add_table(rows=1 + len(rows), cols=len(headers))
    _set_table_style(table)
    _docx_header_row(table, headers)
    _docx_data_rows(table, rows, start_row=1)
    doc.save(str(path))


def write_docx_two_part(path, title, sets_headers, sets_rows, voicing_headers, voicing_rows):
    if not DOCX_AVAILABLE:
        return
    doc = Document()
    doc.core_properties.title = title
    heading = doc.add_heading(title, level=1)
    heading.runs[0].font.size = Pt(14)

    doc.add_heading('Part A: Set-Level Table', level=2)
    tbl_a = doc.add_table(rows=1 + len(sets_rows), cols=len(sets_headers))
    _set_table_style(tbl_a)
    _docx_header_row(tbl_a, sets_headers)
    _docx_data_rows(tbl_a, sets_rows, start_row=1)

    doc.add_page_break()
    doc.add_heading('Part B: Superimposition Table', level=2)
    tbl_b = doc.add_table(rows=1 + len(voicing_rows), cols=len(voicing_headers))
    _set_table_style(tbl_b)
    _docx_header_row(tbl_b, voicing_headers)
    _docx_data_rows(tbl_b, voicing_rows, start_row=1)

    doc.save(str(path))


def write_excel_xml(path, workbook_title, sheets):
    # SpreadsheetML 2003 XML; opens directly in Excel.
    parts = [
        "<?xml version='1.0'?>",
        "<?mso-application progid='Excel.Sheet'?>",
        "<Workbook xmlns='urn:schemas-microsoft-com:office:spreadsheet'"
        " xmlns:o='urn:schemas-microsoft-com:office:office'"
        " xmlns:x='urn:schemas-microsoft-com:office:excel'"
        " xmlns:ss='urn:schemas-microsoft-com:office:spreadsheet'"
        " xmlns:html='http://www.w3.org/TR/REC-html40'>",
        "<DocumentProperties xmlns='urn:schemas-microsoft-com:office:office'>"
        f"<Title>{to_xml(workbook_title)}</Title>"
        "</DocumentProperties>",
    ]

    for sheet_name, headers, rows in sheets:
        parts.append(f"<Worksheet ss:Name='{to_xml(sheet_name)}'><Table>")
        parts.append("<Row>")
        for h in headers:
            parts.append(f"<Cell><Data ss:Type='String'>{to_xml(h)}</Data></Cell>")
        parts.append("</Row>")
        for row in rows:
            parts.append("<Row>")
            for value in row:
                parts.append(f"<Cell><Data ss:Type='String'>{to_xml(value)}</Data></Cell>")
            parts.append("</Row>")
        parts.append("</Table></Worksheet>")

    parts.append("</Workbook>")
    path.write_text("".join(parts), encoding="utf-8")


def export_tables():
    EXPORT_DIR.mkdir(parents=True, exist_ok=True)

    data = json.loads(DATA_PATH.read_text())

    sets_csv = EXPORT_DIR / "appendix_sets_table.csv"
    voicings_csv = EXPORT_DIR / "appendix_superimpositions_table.csv"
    sets_tex = EXPORT_DIR / "appendix_sets_table.tex"
    voicings_tex = EXPORT_DIR / "appendix_superimpositions_table.tex"
    two_part_tex = EXPORT_DIR / "appendix_two_part_tables.tex"

    word_single = EXPORT_DIR / "appendix_single_part_word.doc"
    word_two_part = EXPORT_DIR / "appendix_two_part_word.doc"
    docx_single = EXPORT_DIR / "appendix_single_part.docx"
    docx_two_part = EXPORT_DIR / "appendix_two_part.docx"

    excel_single = EXPORT_DIR / "appendix_single_part_excel.xml"
    excel_two_part = EXPORT_DIR / "appendix_two_part_excel.xml"

    set_headers = [
        "Index",
        "ModernGamut_PCsToC",
        "PitchClasses_TransposedTo0",
        "PrimeForm",
        "ForteNumber",
        "YamaguchiSet",
        "IntervalStructure",
        "ParentPentatonic",
        "DegreesOfAlteration",
    ]

    voicing_headers = [
        "SetIndex",
        "SetPrimeForm",
        "ForteNumber",
        "Bass",
        "ScaleDegree",
        "ChordTones",
        "ChordSymbols",
        "Collections",
    ]

    single_headers = [
        "SetIndex",
        "ModernGamut_PCsToC",
        "PitchClasses_TransposedTo0",
        "PrimeForm",
        "ForteNumber",
        "YamaguchiSet",
        "IntervalStructure",
        "ParentPentatonic",
        "DegreesOfAlteration",
        "Bass",
        "ScaleDegree",
        "ChordTones",
        "ChordSymbols",
        "Collections",
    ]

    sets_rows = []
    voicing_rows = []
    single_rows = []

    with sets_csv.open("w", newline="", encoding="utf-8") as f_set:
        writer = csv.writer(f_set)
        writer.writerow(set_headers)

        with voicings_csv.open("w", newline="", encoding="utf-8") as f_voicing:
            vwriter = csv.writer(f_voicing)
            vwriter.writerow(voicing_headers)

            for idx, item in enumerate(data, 1):
                names = join(item.get("names_transposed_to_C", []))
                pcs0 = join(item.get("pcs_transposed_to_0", []))
                prime = "(" + "".join(str(x) for x in item.get("prime_form", [])) + ")"
                forte = item.get("forte", "")
                yamaguchi = item.get("yamaguchi_set", "")
                interval_struct = item.get("interval_structure", "")
                parent_labels = join(get_parent_pentatonic_labels(item))
                alteration_count = get_alteration_count(item.get("pentatonic_reference") or [])
                voicings = item.get("voicings") or []

                writer.writerow([
                    idx,
                    names,
                    pcs0,
                    prime,
                    forte,
                    yamaguchi,
                    interval_struct,
                    parent_labels,
                    alteration_count,
                ])
                sets_rows.append([
                    idx,
                    names,
                    pcs0,
                    prime,
                    forte,
                    yamaguchi,
                    interval_struct,
                    parent_labels,
                    alteration_count,
                ])

                for voicing in voicings:
                    collections = join(get_voicing_collections(voicing))
                    vwriter.writerow([
                        idx,
                        prime,
                        forte,
                        voicing.get("bass_name", ""),
                        voicing.get("inversion", ""),
                        join(voicing.get("chord_tones") or []),
                        voicing.get("chord_symbol", ""),
                        collections,
                    ])
                    voicing_rows.append([
                        idx,
                        prime,
                        forte,
                        voicing.get("bass_name", ""),
                        voicing.get("inversion", ""),
                        join(voicing.get("chord_tones") or []),
                        voicing.get("chord_symbol", ""),
                        collections,
                    ])
                    single_rows.append([
                        idx,
                        names,
                        pcs0,
                        prime,
                        forte,
                        yamaguchi,
                        interval_struct,
                        parent_labels,
                        alteration_count,
                        voicing.get("bass_name", ""),
                        voicing.get("inversion", ""),
                        join(voicing.get("chord_tones") or []),
                        voicing.get("chord_symbol", ""),
                        collections,
                    ])

    # Dissertation-friendly longtable (set-level)
    with sets_tex.open("w", encoding="utf-8") as f_tex:
        f_tex.write("% Appendix table generated from data/sets.json\n")
        f_tex.write("\\begingroup\\small\n")
        f_tex.write("\\setlength{\\LTleft}{0pt}\n")
        f_tex.write("\\setlength{\\LTright}{0pt}\n")
        f_tex.write("\\begin{longtable}{r p{2.9cm} p{2.2cm} p{1.3cm} p{1.4cm} p{1.6cm} p{2.0cm} p{3.5cm} c}\\n")
        f_tex.write("\\caption{All 12-TET Pentatonic Sets with Classification Metadata}\\\\n")
        f_tex.write("\\hline\n")
        f_tex.write("\\textbf{\\#} & \\textbf{Modern Gamut} & \\textbf{PCs} & \\textbf{Prime} & \\textbf{Forte} & \\textbf{Yamaguchi} & \\textbf{Interval} & \\textbf{Parent Pentatonic} & \\textbf{Alt} \\\\\\n+")
        f_tex.write("\\hline\\endfirsthead\n")
        f_tex.write("\\hline\n")
        f_tex.write("\\textbf{\\#} & \\textbf{Modern Gamut} & \\textbf{PCs} & \\textbf{Prime} & \\textbf{Forte} & \\textbf{Yamaguchi} & \\textbf{Interval} & \\textbf{Parent Pentatonic} & \\textbf{Alt} \\\\\\n+")
        f_tex.write("\\hline\\endhead\n")

        for row in sets_rows:
            line = f"{row[0]} & {to_tex(row[1])} & {to_tex(row[2])} & {to_tex(row[3])} & {to_tex(row[4])} & {to_tex(row[5])} & {to_tex(row[6])} & {to_tex(row[7])} & {row[8]} \\\\n+"
            f_tex.write(line)

        f_tex.write("\\hline\n")
        f_tex.write("\\end{longtable}\n")
        f_tex.write("\\endgroup\n")

    # Two-part LaTeX: superimposition table and wrapper include file
    with voicings_tex.open("w", encoding="utf-8") as f_tex:
        f_tex.write("% Appendix superimposition table generated from data/sets.json\n")
        f_tex.write("\\begingroup\\scriptsize\n")
        f_tex.write("\\setlength{\\LTleft}{0pt}\n")
        f_tex.write("\\setlength{\\LTright}{0pt}\n")
        f_tex.write("\\begin{longtable}{r p{1.2cm} p{1.3cm} p{1.0cm} p{1.2cm} p{2.3cm} p{2.3cm} p{5.0cm}}\n")
        f_tex.write("\\caption{Superimposition Details for All Pentatonic Sets}\\\\\n")
        f_tex.write("\\hline\n")
        f_tex.write("\\textbf{Set} & \\textbf{Prime} & \\textbf{Forte} & \\textbf{Bass} & \\textbf{Degree} & \\textbf{Chord Tones} & \\textbf{Chord Symbols} & \\textbf{Collections} \\\\\n")
        f_tex.write("\\hline\\endfirsthead\n")
        f_tex.write("\\hline\n")
        f_tex.write("\\textbf{Set} & \\textbf{Prime} & \\textbf{Forte} & \\textbf{Bass} & \\textbf{Degree} & \\textbf{Chord Tones} & \\textbf{Chord Symbols} & \\textbf{Collections} \\\\\n")
        f_tex.write("\\hline\\endhead\n")
        for row in voicing_rows:
            f_tex.write(
                f"{row[0]} & {to_tex(row[1])} & {to_tex(row[2])} & {to_tex(row[3])} & {to_tex(row[4])} & {to_tex(row[5])} & {to_tex(row[6])} & {to_tex(row[7])} \\\\\n"
            )
        f_tex.write("\\hline\n")
        f_tex.write("\\end{longtable}\n")
        f_tex.write("\\endgroup\n")

    with two_part_tex.open("w", encoding="utf-8") as f:
        f.write("% Two-part appendix include file\n")
        f.write("% Part A:\n")
        f.write("\\input{appendix_sets_table.tex}\n\n")
        f.write("% Part B:\n")
        f.write("\\input{appendix_superimpositions_table.tex}\n")

    # Word-friendly exports
    write_word_doc(
        word_single,
        "Appendix (Single-Part): Pentatonic Sets with Superimpositions",
        single_headers,
        single_rows,
    )
    write_word_two_part_doc(
        word_two_part,
        "Appendix (Two-Part): Pentatonic Set and Superimposition Tables",
        set_headers,
        sets_rows,
        voicing_headers,
        voicing_rows,
    )

    # Native .docx exports
    write_docx_single(
        docx_single,
        "Appendix (Single-Part): Pentatonic Sets with Superimpositions",
        single_headers,
        single_rows,
    )
    write_docx_two_part(
        docx_two_part,
        "Appendix (Two-Part): Pentatonic Set and Superimposition Tables",
        set_headers,
        sets_rows,
        voicing_headers,
        voicing_rows,
    )

    # Excel-friendly exports
    write_excel_xml(
        excel_single,
        "Appendix Single-Part Table",
        [("SinglePart", single_headers, single_rows)],
    )
    write_excel_xml(
        excel_two_part,
        "Appendix Two-Part Tables",
        [
            ("Sets", set_headers, sets_rows),
            ("Superimpositions", voicing_headers, voicing_rows),
        ],
    )

    print(f"Wrote {sets_csv}")
    print(f"Wrote {voicings_csv}")
    print(f"Wrote {sets_tex}")
    print(f"Wrote {voicings_tex}")
    print(f"Wrote {two_part_tex}")
    print(f"Wrote {word_single}")
    print(f"Wrote {word_two_part}")
    if DOCX_AVAILABLE:
        print(f"Wrote {docx_single}")
        print(f"Wrote {docx_two_part}")
    print(f"Wrote {excel_single}")
    print(f"Wrote {excel_two_part}")


if __name__ == "__main__":
    export_tables()
